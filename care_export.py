"""Printable care export; generated in memory and never published by URL."""
import io
import re
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


def build_care_docx(pet, entries, courses, doses, operations, actors, labels, t=lambda s: s, scope='all'):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 2'].font.size = Pt(12)

    def text(value):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value if value is not None else ''))
    def para(value, style=None):
        return doc.add_paragraph(text(value), style)
    def field(label, value):
        p = para('')
        p.add_run(t(label) + ': ').bold = True
        p.add_run(text(value))
    def local(value, zone):
        return value.replace(tzinfo=timezone.utc).astimezone(ZoneInfo(zone)).strftime('%d.%m.%Y %H:%M')
    def table(headers, rows, widths):
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.autofit = False
        for column, width in zip(tb.columns, widths):
            column.width = Cm(width)
        for cell, label in zip(tb.rows[0].cells, headers):
            cell.text = t(label)
            for run in cell.paragraphs[0].runs: run.bold = True
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF2')
            cell._tc.get_or_add_tcPr().append(shade)
        repeat = OxmlElement('w:tblHeader')
        tb.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for cell, value in zip(tb.add_row().cells, values): cell.text = text(value)
        for row in tb.rows:
            no_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(no_split)
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ('top', 'left', 'bottom', 'right'):
                    line = OxmlElement('w:' + edge)
                    for key, value in [('val','single'),('sz','4'),('color','D9D9D9')]:
                        line.set(qn('w:' + key), value)
                    borders.append(line)
                props.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    node = OxmlElement('w:' + edge)
                    node.set(qn('w:w'), '90'); node.set(qn('w:type'), 'dxa')
                    margins.append(node)
                props.append(margins)
        para('')
        return tb

    titles = {'all':'Дневник ухода за питомцем', 'postop':'Послеоперационный дневник',
              'chronic':'Дневник наблюдений', 'medications':'Дневник назначений'}
    para(t(titles[scope]), 'Title')
    para(text(pet.name), 'Heading 1')
    field('Питомец', ' · '.join(text(v) for v in (
        t(pet.species), pet.breed, pet.age,
        str(pet.weight_kg) + ' ' + t('кг') if pet.weight_kg is not None else '') if v))
    field('Дата выгрузки', datetime.now(timezone.utc).strftime('%d.%m.%Y %H:%M UTC'))
    para(t('Записи владельца о состоянии питомца и выполнении назначений. Время указано с часовым поясом каждой записи.'))
    if scope != 'medications':
        doc.add_heading(t('Наблюдения'), 1)
        if not entries: para(t('Записей пока нет.'))
        for entry in entries:
            doc.add_heading(local(entry.observed_at, entry.timezone) + ' · ' + entry.timezone, 2)
            field('Вид дневника', t('После операции' if entry.kind == 'postop' else 'Хроническое заболевание'))
            if entry.condition: field('Заболевание или симптом', entry.condition)
            operation = operations.get(entry.operation_id)
            if operation:
                field('Операция', operation.title + (' · ' + operation.operation_date.strftime('%d.%m.%Y') if operation.operation_date else ''))
            field('Добавил(а)', actors.get(entry.created_by, '—'))
            pairs = []
            for key, label in [('appetite','Аппетит'),('activity','Активность'),('pain','Боль')]:
                pairs.append((t(label), t(labels[key].get(getattr(entry, key), '—'))))
            if entry.kind == 'postop': pairs.append((t('Шов'), t(labels['wound'].get(entry.wound, '—'))))
            for key, label in [('weight','Вес, кг'),('vomiting','Рвота за день, эпизодов'),('seizures','Приступы за день, эпизодов')]:
                value = getattr(entry, key)
                pairs.append((t(label), t('Не измеряли') if value is None else str(value)))
            table(['Показатель','Наблюдение'], pairs, [6,11])
            if entry.flags: field('Тревожные признаки', '; '.join(t(labels['flags'].get(f, f)) for f in entry.flags))
            if entry.notes: field('Заметка', entry.notes)
            if entry.photo:
                try:
                    with Image.open(io.BytesIO(entry.photo)) as im:
                        width, height = im.size
                    scale = min(12 / width, 9 / height)
                    caption = para(t('Фото к наблюдению') + ' · ' + local(entry.observed_at, entry.timezone))
                    caption.paragraph_format.keep_with_next = True
                    doc.add_picture(io.BytesIO(entry.photo), width=Cm(width * scale), height=Cm(height * scale))
                except (OSError, ValueError):
                    para(t('Фото не удалось включить в файл. Оно доступно в дневнике на сайте.'))
    if scope in ('all', 'medications'):
        doc.add_heading(t('Назначения и отметки о приёме'), 1)
        if not courses: para(t('Назначений пока нет.'))
        by_course = {}
        for dose in doses: by_course.setdefault(dose.course_id, []).append(dose)
        for course in courses:
            doc.add_heading(text(course.name), 2)
            field('Доза и способ введения из назначения', course.dose)
            if course.instructions: field('Указания врача', course.instructions)
            field('Курс', course.start_date.strftime('%d.%m.%Y') + ' – ' + course.end_date.strftime('%d.%m.%Y'))
            field('Время каждого приёма', ', '.join(course.times) + ' · ' + course.timezone)
            field('Статус', t('Активен' if course.active else 'Остановлен'))
            field('Добавил(а)', actors.get(course.created_by, '—'))
            rows = []
            for dose in by_course.get(course.id, []):
                state = t(labels['states'].get(dose.state, dose.state))
                if dose.state == 'pending' and dose.due_at > datetime.utcnow():
                    state = t('Предстоит')
                mark = (local(dose.marked_at, course.timezone) + '\n' + actors.get(dose.marked_by, '—')) if dose.marked_at else '—'
                rows.append((local(dose.due_at, course.timezone), state, mark))
            if rows: table(['По расписанию', 'Отметка', 'Когда и кем отмечено'], rows, [5,5,7])
            else: para(t('Отметок пока нет.'))
    doc.core_properties.title = t(titles[scope])
    doc.core_properties.author = 'МойДоктор'
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
