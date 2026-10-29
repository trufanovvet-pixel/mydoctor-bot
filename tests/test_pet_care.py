import io
import json
import re
import secrets
from datetime import datetime, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
import pytest
from sqlalchemy import select, func
from werkzeug.datastructures import MultiDict
from PIL import Image
import storage
import pet_care as care
import care_notifications as notices
from care_models import *
from patient_records import OperationRecord
from test_web_app import reset_db, register, web_app


@pytest.fixture(autouse=True)
def clean(): reset_db()


def account(email='care@example.com'):
    c=web_app.app.test_client(); register(c,email)
    with c.session_transaction() as s: uid=s['uid']
    c.post('/pets',data={'name':'Archie','species':'Собака','weight':'12'})
    with storage.SessionLocal() as db: pid=db.scalar(select(storage.Pet.id).where(storage.Pet.user_id==uid))
    assert c.get('/care').status_code==200
    return c,uid,pid


def post(c,path,data=None,**kwargs):
    with c.session_transaction() as s: token=s['care_csrf']
    values=MultiDict()
    for k,v in (data or {}).items(): values[k]=v
    values['care_csrf']=token
    return c.post(path,data=values,**kwargs)


def medication_form(**changes):
    day=(datetime.utcnow()+timedelta(days=1)).date().isoformat()
    return dict(request_key=secrets.token_hex(16),name='Prescribed medicine',dose='As prescribed',instructions='With food',
        confirmed='1',start_date=day,end_date=day,times='08:00, 20:00',timezone='Europe/Moscow',**changes)


def observation_form():
    return dict(request_key=secrets.token_hex(16),kind='chronic',condition='Epilepsy',observed_at=datetime.utcnow().strftime('%Y-%m-%dT%H:%M'),
        timezone='UTC',appetite='normal',activity='normal',pain='none',wound='not_applicable',notes='',vomiting='0',seizures='0')


def due_course(uid,pid,delta=-1):
    now=datetime.utcnow()
    with storage.SessionLocal() as db:
        course=MedicationCourse(pet_id=pid,created_by=uid,request_key=secrets.token_hex(16),name='Test medication',dose='Test prescription',
            timezone='Asia/Makassar',times=['20:00'],start_date=now.date(),end_date=now.date())
        db.add(course);db.flush();dose=MedicationDose(course_id=course.id,due_at=now+timedelta(minutes=delta));db.add(dose);db.commit()
        return course.id,dose.id


def invite(a,pid,b):
    response=post(a,f'/care/pets/{pid}/invites')
    assert response.status_code==200
    path=re.search(r'value="(/care/join/[^"]+)"',response.get_data(as_text=True)).group(1)
    assert b.get(path).status_code==200
    assert post(b,path).status_code==303
    return path


def picture():
    data=io.BytesIO();Image.new('RGB',(24,24),'blue').save(data,format='PNG');data.seek(0);return data


def test_navigation_csrf_and_owner_isolation():
    a,uid,pid=account();b,_,_=account('other@example.com')
    assert a.get(f'/care/pets/{pid}').status_code==200
    assert 'private, no-store' in a.get(f'/care/pets/{pid}').headers['Cache-Control']
    assert b.get(f'/care/pets/{pid}').status_code==404
    assert post(b,f'/care/pets/{pid}/medications',medication_form()).status_code==404
    assert a.post(f'/care/pets/{pid}/medications',data=medication_form()).status_code==400
    assert web_app.app.test_client().get('/care').status_code==302


def test_schedule_confirmation_idempotency_and_free_manual_entry(monkeypatch):
    a,uid,pid=account();values=medication_form()
    storage.set_bot_setting('billing_live','1')
    from test_clinical_safety_access import exhaust
    exhaust(uid)
    response = post(a,f'/care/pets/{pid}/medications',{**values,'confirmed':'0'})
    assert response.status_code == 200
    assert 'data-medication-error' in response.text
    with storage.SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(MedicationCourse)) == 0
        assert db.scalar(select(func.count()).select_from(MedicationDose)) == 0
    for _ in range(2): assert post(a,f'/care/pets/{pid}/medications',values).status_code==303
    with storage.SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(MedicationCourse))==1
        assert db.scalar(select(func.count()).select_from(MedicationDose))==2
        course=db.scalar(select(MedicationCourse));dose=db.scalar(select(MedicationDose).order_by(MedicationDose.due_at))
        assert dose.due_at.hour==5 and course.dose=='As prescribed'
        from billing import CreditUsage
        assert db.scalar(select(func.count()).select_from(CreditUsage))==0


@pytest.mark.parametrize('values',[
    {'times':'25:00'}, {'times':'08:00, 08:00'}, {'timezone':'wrong/timezone'},
    {'end_date':'2026-02-30'}, {'start_date':'2020-01-01'}, {'times':''}, {'times':'08:00,'*7},
])
def test_invalid_schedule_rejected(values):
    form=medication_form();form.update(values)
    with pytest.raises(ValueError): care.schedule(form)


def test_dst_and_timezone_are_explicit_and_never_double_a_slot():
    form=medication_form();form.update(start_date='2026-11-01',end_date='2026-11-01',times='01:30',timezone='America/New_York')
    with pytest.raises(ValueError,match='перевод часов'): care.schedule(form,now=datetime(2026,10,27))
    form.update(start_date='2026-03-08',end_date='2026-03-08',times='02:30')
    with pytest.raises(ValueError): care.schedule(form,now=datetime(2026,3,1))
    form.update(start_date='2026-10-28',end_date='2026-10-28',times='08:00, 20:00',timezone='Asia/Makassar')
    _,slots=care.schedule(form,now=datetime(2026,10,27));assert slots==[datetime(2026,10,28),datetime(2026,10,28,12)]


def test_family_single_use_link_shared_dose_and_revocation():
    a,uid,pid=account();b,buid,_=account('family@example.com');outsider,_,_=account('out@example.com')
    cid,did=due_course(uid,pid)
    path=invite(a,pid,b)
    assert outsider.get(path).status_code==410
    assert post(b,f'/care/pets/{pid}/invites').status_code==404
    assert post(b,f'/care/doses/{did}',{'state':'taken'}).status_code==303
    assert post(a,f'/care/doses/{did}',{'state':'taken'}).status_code==409
    with storage.SessionLocal() as db:
        dose=db.get(MedicationDose,did);assert dose.marked_by==buid and dose.state=='taken'
        inv=db.scalar(select(CareInvite));assert inv.digest != path.rsplit('/',1)[1]
    assert post(a,f'/care/pets/{pid}/members/{buid}/remove').status_code==303
    assert b.get(f'/care/pets/{pid}').status_code==404
    assert post(b,f'/care/medications/{cid}/stop').status_code==404
    with storage.SessionLocal() as db: assert db.get(MedicationDose,did).state=='taken'


def test_invite_preview_does_not_consume_and_revocation_invalidates_link():
    a,uid,pid=account();b,_,_=account('family@example.com')
    response=post(a,f'/care/pets/{pid}/invites');path=re.search(r'value="(/care/join/[^"]+)"',response.text).group(1)
    for _ in range(2): assert b.get(path).status_code==200
    assert post(a,f'/care/pets/{pid}/invites/revoke').status_code==303
    assert post(b,path).status_code==410


def test_future_dose_cannot_be_marked_and_stop_cancels_only_pending():
    a,uid,pid=account();cid,did=due_course(uid,pid,delta=30)
    assert post(a,f'/care/doses/{did}',{'state':'taken'}).status_code==409
    assert post(a,f'/care/medications/{cid}/stop').status_code==303
    with storage.SessionLocal() as db: assert db.get(MedicationDose,did).state=='cancelled'


def test_duplicate_active_schedule_and_cross_pet_status_are_blocked():
    a,uid,pid=account();b,_,bpid=account('other@example.com')
    form=medication_form();form['enable_reminders']='1'
    assert post(a,f'/care/pets/{pid}/medications',form).status_code==303
    response = post(a,f'/care/pets/{pid}/medications',medication_form())
    assert response.status_code == 200
    assert 'data-medication-error' in response.text
    with storage.SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(MedicationCourse)) == 1
        assert db.scalar(select(func.count()).select_from(MedicationDose)) == 2
    with storage.SessionLocal() as db:
        assert db.get(CarePreference,(pid,uid)).reminders
        did=db.scalar(select(MedicationDose.id))
    assert b.get(f'/care/pets/{pid}/dose-status?id={did}').status_code==404
    assert b.get(f'/care/pets/{bpid}/dose-status?id={did}').json['doses']==[]


def test_family_poll_updates_actor_and_does_not_expose_billing_or_private_files():
    a,uid,pid=account();b,buid,_=account('family@example.com');invite(a,pid,b)
    cid,did=due_course(uid,pid);post(b,f'/care/doses/{did}',{'state':'taken'})
    data=a.get(f'/care/pets/{pid}/dose-status?id={did}').json['doses'][0]
    assert data['state']=='taken' and data['actor']=='QA'
    with storage.SessionLocal() as db:
        doc=web_app.WebDocument(user_id=uid,pet_id=pid,filename='private.pdf',mime_type='application/pdf',data=b'%PDF');db.add(doc);db.commit();fid=doc.id
    assert b.get(f'/documents/{fid}').status_code==404
    assert b.get(f'/history?pet_id={pid}').status_code==404


def test_expired_invite_cannot_add_member():
    a,uid,pid=account();b,buid,_=account('family@example.com')
    response=post(a,f'/care/pets/{pid}/invites');path=re.search(r'value="(/care/join/[^"]+)"',response.text).group(1)
    with storage.SessionLocal() as db:
        db.scalar(select(CareInvite)).expires_at=datetime.utcnow()-timedelta(seconds=1);db.commit()
    assert post(b,path).status_code==410
    with storage.SessionLocal() as db: assert db.get(CareMember,(pid,buid)) is None


def test_owner_cannot_attach_another_pets_operation():
    a,uid,pid=account();b,buid,bpid=account('other@example.com')
    with storage.SessionLocal() as db:
        op=OperationRecord(user_id=buid,pet_id=bpid,title='Private operation');db.add(op);db.commit();oid=op.id
    values=observation_form();values.update(kind='postop',operation_id=oid)
    assert post(a,f'/care/pets/{pid}/observations',values).status_code==400


def test_observation_photo_family_access_and_urgent_guidance():
    a,uid,pid=account();b,buid,_=account('family@example.com');c,_,_=account('third@example.com');invite(a,pid,b)
    with storage.SessionLocal() as db:
        op=OperationRecord(user_id=uid,pet_id=pid,title='TPLO',operation_date=datetime.utcnow().date());db.add(op);db.commit();oid=op.id
    values=observation_form();values.update(kind='postop',operation_id=str(oid),flags='breathing',photo=(picture(),'incision.png'))
    r=post(b,f'/care/pets/{pid}/observations',values,content_type='multipart/form-data');assert r.status_code==303
    with storage.SessionLocal() as db:
        entry=db.scalar(select(HealthEntry));eid=entry.id;assert entry.urgency=='urgent' and entry.created_by==buid
        assert entry.photo[:2]==b'\xff\xd8'
    for client in (a,b): assert client.get(f'/care/photos/{eid}').status_code==200
    assert c.get(f'/care/photos/{eid}').status_code==404
    assert 'выезжайте сейчас' in a.get(f'/care/pets/{pid}').text
    post(a,f'/care/pets/{pid}/members/{buid}/remove')
    assert b.get(f'/care/photos/{eid}').status_code==404


@pytest.mark.parametrize('changes,expected',[
    ({},'recorded'), ({'wound':'worse'},'contact'), ({'wound':'open'},'contact'),
    ({'appetite':'none'},'contact'), ({'activity':'very_low'},'contact'), ({'pain':'severe'},'contact'),
    ({'flags':['collapse']},'urgent'), ({'flags':['urine']},'urgent'), ({'seizures':'1'},'contact'),
])
def test_reported_signs_drive_urgency(changes,expected):
    form=MultiDict(observation_form());form.update(changes)
    # update on MultiDict appends values; assign explicitly to simulate a fresh form.
    for k,v in changes.items():
        if isinstance(v,list): form.setlist(k,v)
        else: form[k]=v
    assert care.observation(form)['urgency']==expected


@pytest.mark.parametrize('field,value',[('weight','NaN'),('weight','-2'),('vomiting','-1'),('seizures','1.2'),('observed_at','2040-01-01T10:00')])
def test_invalid_observations_are_not_saved(field,value):
    a,_,pid=account();form=observation_form();form[field]=value
    assert post(a,f'/care/pets/{pid}/observations',form).status_code==400
    with storage.SessionLocal() as db: assert db.scalar(select(func.count()).select_from(HealthEntry))==0


def test_chronic_empty_measurements_are_not_zero_and_entries_idempotent():
    a,uid,pid=account();form=observation_form();form.update(weight='',vomiting='',seizures='')
    for _ in range(2): assert post(a,f'/care/pets/{pid}/observations',form).status_code==303
    with storage.SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(HealthEntry))==1
        entry=db.scalar(select(HealthEntry));assert entry.weight is None and entry.vomiting is None


def test_doctor_requires_explicit_share_and_prior_consultation():
    from test_doctor_portal import doctor
    a,uid,pid=account();d=doctor();path=f'/doctor/care/pets/{pid}'
    assert d.get(path).status_code==404
    post(a,f'/care/pets/{pid}/preferences',{'doctor_enabled':'1'})
    assert d.get(path).status_code==404
    with storage.SessionLocal() as db:
        db.add(storage.Consultation(user_id=uid,pet_id=pid,kind='consult_request',user_text='test',assistant_text=''));db.commit()
    assert d.get(path).status_code==200
    assert 'name="confirmed"' not in d.get(path).text
    post(a,f'/care/pets/{pid}/preferences',{})
    assert d.get(path).status_code==404


def test_ocr_does_not_schedule_without_review_and_failure_refunds(monkeypatch):
    a,uid,pid=account()
    response=SimpleNamespace(output_text=json.dumps({'medicines':[{'name':'Original name','dose':'Original dose','instructions':'Original directions','times':''}]}))
    with patch.object(web_app.client.responses,'create',return_value=response):
        r=post(a,f'/care/pets/{pid}/prescriptions',{'request_key':secrets.token_hex(16),'file':(picture(),'prescription.png')},content_type='multipart/form-data')
    assert r.status_code==303
    with storage.SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(MedicationCourse))==0
        source=db.scalar(select(CarePrescription));assert source.medicines[0]['times']=='';did=source.document_id
    html=a.get(r.headers['Location']).text;assert 'Original name' in html and 'Черновик' in html
    assert a.get(f'/care/prescriptions/{did}').status_code==200
    b,_,_=account('outside@example.com');assert b.get(f'/care/prescriptions/{did}').status_code==404
    with patch.object(web_app.client.responses,'create',side_effect=RuntimeError('offline')):
        r=post(a,f'/care/pets/{pid}/prescriptions',{'request_key':secrets.token_hex(16),'file':(picture(),'prescription.png')},content_type='multipart/form-data')
    assert r.status_code==422
    with storage.SessionLocal() as db:
        from billing import CreditUsage
        assert db.scalar(select(CreditUsage).order_by(CreditUsage.created_at.desc())).status=='failed'


def test_reminders_require_opt_in_and_deduplicate_and_skip_marked_doses():
    a,uid,pid=account();cid,did=due_course(uid,pid)
    with storage.SessionLocal() as db: db.get(storage.User,uid).telegram_id=77553311;db.commit()
    notices.enqueue_due()
    with storage.SessionLocal() as db: assert db.scalar(select(func.count()).select_from(CareNotice))==0
    post(a,f'/care/pets/{pid}/preferences',{'reminders':'1'})
    notices.enqueue_due();notices.enqueue_due()
    with storage.SessionLocal() as db: assert db.scalar(select(func.count()).select_from(CareNotice))==1
    nid=notices.claim_notice();assert nid is not None
    assert notices.claim_notice() is None
    assert notices.notice_payload(nid)['dose_id']==did
    post(a,f'/care/doses/{did}',{'state':'taken'})
    assert notices.notice_payload(nid) is None
    with storage.SessionLocal() as db: assert db.get(CareNotice,nid).state=='skipped'


def test_revocation_or_opt_out_blocks_queued_reminder():
    a,uid,pid=account();b,buid,_=account('family@example.com');invite(a,pid,b);due_course(uid,pid)
    with storage.SessionLocal() as db: db.get(storage.User,buid).telegram_id=99112233;db.commit()
    post(b,f'/care/pets/{pid}/preferences',{'reminders':'1'});notices.enqueue_due();nid=notices.claim_notice()
    assert notices.notice_payload(nid)
    post(a,f'/care/pets/{pid}/members/{buid}/remove');assert notices.notice_payload(nid) is None


@pytest.mark.asyncio
async def test_telegram_reminder_is_retried_after_transport_failure():
    a,uid,pid=account();due_course(uid,pid)
    with storage.SessionLocal() as db: db.get(storage.User,uid).telegram_id=11229933;db.commit()
    post(a,f'/care/pets/{pid}/preferences',{'reminders':'1'})
    application=SimpleNamespace(bot=SimpleNamespace(send_message=AsyncMock(side_effect=TimeoutError())))
    await notices.deliver_care_notifications(application)
    with storage.SessionLocal() as db:
        item=db.scalar(select(CareNotice));assert item.state=='retry' and item.attempts==1
        item.next_attempt=datetime.utcnow()-timedelta(seconds=1);db.commit()
    application.bot.send_message=AsyncMock()
    await notices.deliver_care_notifications(application)
    assert application.bot.send_message.await_count==1
    with storage.SessionLocal() as db: assert db.scalar(select(CareNotice)).state=='delivered'
    await notices.deliver_care_notifications(application)
    assert application.bot.send_message.await_count==1


def test_stale_reminders_and_changed_push_binding_are_skipped():
    from web_push import PushSubscription
    a,uid,pid=account();cid,did=due_course(uid,pid)
    post(a,f'/care/pets/{pid}/preferences',{'reminders':'1'})
    with storage.SessionLocal() as db:
        db.add(PushSubscription(id='x'*64,role='owner',subject_id=uid,endpoint='https://fcm.googleapis.com/fcm/send/test',
            p256dh='key',auth='auth',binding='old',expires_at=datetime.utcnow()+timedelta(days=1)));db.commit()
    notices.enqueue_due();nid=notices.claim_notice();assert notices.notice_payload(nid)['channel']=='push'
    with storage.SessionLocal() as db:
        db.get(PushSubscription,'x'*64).binding='new';db.commit()
    assert notices.notice_payload(nid) is None


def test_english_care_pages_and_forms_preserve_stored_values():
    from test_owner_language import VisibleText
    a,uid,pid=account();a.get('/language/en')
    for path in ['/care',f'/care/pets/{pid}']:
        response=a.get(path);assert response.status_code==200
        visible=VisibleText();visible.feed(response.text)
        assert not re.search('[А-Яа-яЁё]',''.join(visible.text)), ''.join(visible.text)


@pytest.mark.asyncio
async def test_telegram_care_entrypoint_uses_private_login(bot,ctx,update):
    event=update(notices.LABEL)
    await bot.message(event,ctx)
    url=event.effective_message.reply_text.call_args.kwargs['reply_markup'].inline_keyboard[0][0].url
    assert '?next=/care' in url and '/login/' in url
    c=web_app.app.test_client();page=c.get(url.split('https://mydoctor.vet')[1]);assert page.status_code==200
    csrf=re.search('name="login_csrf" value="([^"]+)"',page.text).group(1)
    r=c.post(url.split('https://mydoctor.vet')[1],data={'login_csrf':csrf,'next':'/care'})
    assert r.status_code==303 and r.headers['Location']=='/care'


def test_word_export_full_history_photos_statuses_and_access():
    from docx import Document
    from zipfile import ZipFile
    a,uid,pid=account()
    b,buid,bpid=account('export-outsider@example.com')
    cid,did=due_course(uid,pid)
    assert post(a,f'/care/doses/{did}',{'state':'taken'}).status_code==303
    with storage.SessionLocal() as db:
        operation=OperationRecord(user_id=uid,pet_id=pid,title='Export operation')
        db.add(operation);db.flush()
        for index in range(35):
            db.add(HealthEntry(pet_id=pid,created_by=uid,request_key=secrets.token_hex(16),
                kind='postop' if index == 0 else 'chronic',operation_id=operation.id if index == 0 else None,
                condition='Export condition',observed_at=datetime.utcnow()-timedelta(days=35-index),
                timezone='Asia/Makassar',appetite='normal',activity='reduced',pain='none',wound='unchanged',
                notes=f'Export entry {index:02}',weight=None,vomiting=0,seizures=None,
                photo=picture().getvalue() if index == 0 else None))
        db.add(HealthEntry(pet_id=bpid,created_by=buid,request_key=secrets.token_hex(16),
            kind='chronic',observed_at=datetime.utcnow(),timezone='UTC',appetite='normal',
            activity='normal',pain='none',notes='OUTSIDER PRIVATE ENTRY'))
        for state in ('skipped','pending','cancelled'):
            db.add(MedicationDose(course_id=cid,due_at=datetime.utcnow()-timedelta(hours=len(state)),state=state))
        db.add(MedicationDose(course_id=cid,due_at=datetime.utcnow()+timedelta(days=1),state='pending'))
        db.commit()
    path=f'/care/pets/{pid}/export.docx'
    response=a.get(path)
    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert 'attachment;' in response.headers['Content-Disposition']
    assert 'no-store' in response.headers['Cache-Control']
    with ZipFile(io.BytesIO(response.data)) as z:
        xml=z.read('word/document.xml').decode()
        assert any(name.startswith('word/media/') for name in z.namelist())
    Document(io.BytesIO(response.data))
    for index in range(35): assert f'Export entry {index:02}' in xml
    for value in ('Export operation','Asia/Makassar','Test prescription','Дал(а)','Пропустил(а)',
                  'Нет отметки','Курс остановлен','Предстоит','Не измеряли'):
        assert value in xml
    assert 'OUTSIDER PRIVATE ENTRY' not in xml
    assert b.get(path).status_code == 404
    assert web_app.app.test_client().get(path).status_code == 302
    assert a.get(path+'?scope=invalid').status_code == 400
    for scope, included, excluded in [
        ('postop','Export entry 00','Export entry 34'),
        ('chronic','Export entry 34','Export entry 00'),
        ('medications','Test prescription','Export entry 00'),
    ]:
        with ZipFile(io.BytesIO(a.get(path+'?scope='+scope).data)) as z:
            part=z.read('word/document.xml').decode()
        assert included in part and excluded not in part
        if scope != 'medications': assert 'Test prescription' not in part
    a.get('/language/en')
    with ZipFile(io.BytesIO(a.get(path+'?scope=medications').data)) as z:
        english=z.read('word/document.xml').decode()
    assert 'Prescription diary' in english and 'Дневник назначений' not in english


def test_word_export_empty_and_doctor_permission_revocation():
    from docx import Document
    from test_doctor_portal import doctor
    a,uid,pid=account()
    response=a.get(f'/care/pets/{pid}/export.docx')
    doc=Document(io.BytesIO(response.data))
    assert 'Записей пока нет.' in [p.text for p in doc.paragraphs]
    d=doctor();path=f'/doctor/care/pets/{pid}/export.docx'
    assert d.get(path).status_code == 404
    post(a,f'/care/pets/{pid}/preferences',{'doctor_enabled':'1'})
    assert d.get(path).status_code == 404
    with storage.SessionLocal() as db:
        db.add(storage.Consultation(user_id=uid,pet_id=pid,kind='consult_request',user_text='test',assistant_text=''))
        db.commit()
    assert d.get(path).status_code == 200
    post(a,f'/care/pets/{pid}/preferences',{})
    assert d.get(path).status_code == 404
